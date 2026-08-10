import json
from datetime import datetime, timedelta
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from src.api import create_app, trading_state
from src.api.trial_run import trial_run_state
from src.strategy import Direction, OrderStatus, Position, Trade
from src.trading.types import MarketData

from tests.test_trial_run_api import _config_path, _trial_config, install_gateway, login
from tests.test_trial_run_closed_loop import _configure, _tick


def _doc_texts(content: bytes) -> list[str]:
    document = Document(BytesIO(content))
    texts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.append(cell.text)
    return texts


def teardown_function():
    trading_state.clear_main()
    trial_run_state.reset()


def test_running_run_cannot_export_report(monkeypatch, tmp_path):
    config_path = _trial_config(_config_path(tmp_path, "report-running"), hold_bars=1)
    _configure(config_path, simulate=False, hold_bars=1)
    monkeypatch.setenv("QUANT_TRIAL_CONFIG", str(config_path))
    install_gateway(monkeypatch)
    app = create_app()

    with TestClient(app) as client:
        login(client)
        assert client.post("/trial-run/prepare").status_code == 200
        response = client.get("/trial-run/report.docx")

    assert response.status_code == 409
    assert response.json()["detail"] == "trial run has no terminal acceptance result"


def test_passed_simulated_report_is_truthful(monkeypatch, tmp_path):
    config_path = _trial_config(_config_path(tmp_path, "report-simulated"), hold_bars=1)
    _configure(config_path, simulate=True, hold_bars=1)
    monkeypatch.setenv("QUANT_TRIAL_CONFIG", str(config_path))
    gateway = install_gateway(monkeypatch)
    app = create_app()
    start = datetime.now()

    with TestClient(app) as client:
        login(client)
        assert client.post("/trial-run/prepare").status_code == 200
        entry = trading_state.get("verify_trial")
        entry.engine.on_tick(_tick(3130, start))
        source = gateway.orders["ORDER_1"]
        source.status = OrderStatus.SUBMITTED
        entry.engine._on_order(source)
        entry.strategy._order_ownership["ORDER_1"]["submitted_monotonic"] -= 3
        assert client.post(
            "/trial-run/simulation/prepare",
            json={"source_order_id": "ORDER_1"},
        ).status_code == 200
        gateway.on_order(source)
        ready = client.post(
            "/trial-run/simulation/prepare",
            json={"source_order_id": "ORDER_1"},
        )
        synthetic_entry_id = ready.json()["status"]["current_order_id"]
        assert client.post(
            "/trial-run/simulate-fill",
            json={"order_id": synthetic_entry_id},
        ).status_code == 200
        entry.engine.on_tick(_tick(3131, start + timedelta(minutes=1)))
        synthetic_close_id = entry.engine.trial_run_execution.current_order_id
        assert client.post(
            "/trial-run/simulate-fill",
            json={"order_id": synthetic_close_id},
        ).status_code == 200
        assert client.get("/trial-run/status").json()["outcome"] == "passed_simulated"

        response = client.get("/trial-run/report.docx")

    assert response.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in response.headers["content-type"]
    assert "trial-run-report" in response.headers["content-disposition"]
    texts = " ".join(_doc_texts(response.content))
    assert "真实报单链路通过" in texts
    assert "模拟成交" in texts
    assert "真实成交通过" not in texts


def test_passed_real_report_is_truthful(monkeypatch, tmp_path):
    config_path = _trial_config(_config_path(tmp_path, "report-real"), hold_bars=1)
    _configure(config_path, simulate=False, hold_bars=1)
    monkeypatch.setenv("QUANT_TRIAL_CONFIG", str(config_path))
    gateway = install_gateway(monkeypatch)
    app = create_app()
    start = datetime.now()

    with TestClient(app) as client:
        login(client)
        assert client.post("/trial-run/prepare").status_code == 200
        entry = trading_state.get("verify_trial")
        entry.engine.on_tick(_tick(3130, start))
        open_order = gateway.orders["ORDER_1"]
        open_order.status = OrderStatus.SUBMITTED
        entry.engine._on_order(open_order)
        gateway.on_position(Position("rb2510", Direction.LONG, 1, price=3130, cost=3130))
        gateway.on_trade(Trade("T-ENTRY", "ORDER_1", "rb2510", Direction.LONG, 3130, 1))
        entry.engine.on_tick(_tick(3131, start + timedelta(minutes=1)))
        close_order = gateway.orders["ORDER_2"]
        close_order.status = OrderStatus.SUBMITTED
        entry.engine._on_order(close_order)
        gateway.on_position(Position("rb2510", Direction.LONG, 0, price=3131, cost=3131))
        gateway.on_trade(Trade("T-CLOSE", "ORDER_2", "rb2510", Direction.SHORT, 3131, 1))
        entry.engine.on_timer(entry.engine._monotonic())
        assert client.get("/trial-run/status").json()["outcome"] == "passed_real"

        response = client.get("/trial-run/report.docx")

    assert response.status_code == 200
    texts = " ".join(_doc_texts(response.content))
    assert "真实成交闭环通过" in texts