"""Generate a truthful DOCX acceptance report for a terminal trial run."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any, Mapping

from docx import Document
from docx.shared import Pt


def _text(value: Any, fallback: str = "--") -> str:
    if value is None:
        return fallback
    text = str(value).strip()
    return text or fallback


def _direction(value: Any) -> str:
    code = str(value or "").strip().lower()
    return {"long": "多", "short": "空"}.get(code, _text(value))


def _offset(value: Any) -> str:
    code = str(value or "").strip().lower()
    return {
        "open": "开仓",
        "close": "平仓",
        "close_today": "平今",
        "close_yesterday": "平昨",
    }.get(code, _text(value))


def _status_label(value: Any) -> str:
    code = str(value or "").strip().lower()
    return {
        "submitting": "提交中",
        "submitted": "已报未成",
        "partfilled": "部分成交",
        "filled": "全部成交",
        "cancelled": "已撤单",
        "canceled": "已撤单",
        "rejected": "已拒单",
    }.get(code, _text(value))


def _format_time(value: Any) -> str:
    text = _text(value, "")
    if not text or text == "--":
        return "--"
    return str(text).replace("T", " ")[:19]


def _conclusion(status: Mapping[str, Any]) -> str:
    outcome = str(status.get("outcome") or "")
    if outcome == "passed_real":
        return "真实成交闭环通过"
    if outcome == "passed_simulated":
        return "真实报单链路通过；成交后处理由模拟成交验证，未取得券商真实成交"
    if outcome == "failed":
        return "试运行失败"
    if outcome == "aborted":
        return "试运行中止，无通过结论"
    return "试运行未产生终态结论"


def generate_trial_run_report(
    status: Mapping[str, Any],
    config: Mapping[str, Any],
) -> bytes:
    """Build a DOCX report from the terminal status and safe config snapshot."""
    document = Document()
    document.add_heading("期货程序化交易试运行验收测试报告", level=0)

    document.add_heading("一、测试基本信息", level=1)
    basic = [
        ("运行 ID", _text(status.get("run_id"))),
        ("生成时间", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")),
        ("账号", _text(config.get("masked_account_id"))),
        ("环境", _text(status.get("runtime_environment") or config.get("environment"))),
        ("固定合约", _text(status.get("symbol") or config.get("allowed_symbol"))),
        ("方向/开平", f"{_direction(status.get('current_track'))} 1 手验证开仓 -> 平仓"),
        ("最终结论", _conclusion(status)),
        ("success_basis", _text(status.get("success_basis"))),
    ]
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in basic:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = value

    document.add_heading("二、前置检查结果", level=1)
    preflight = [
        ("配置有效", "是" if bool(status.get("config_valid")) else "否"),
        ("交易网关连接", "是" if bool(status.get("gateway_connected")) else "否"),
        ("目标合约持仓归零", "是" if int(status.get("broker_position_volume") or 0) == 0 else "否"),
        ("目标合约无活动委托", "是" if not status.get("broker_active_order_ids") else "否"),
        ("模拟成交允许", "是" if bool(status.get("simulate_fill_enabled")) else "否"),
    ]
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in preflight:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = value

    document.add_heading("三、行情证据", level=1)
    market_rows = [
        ("最近有效 tick", _text(status.get("last_market_price"))),
        ("行情时间", _format_time(status.get("last_market_timestamp"))),
        ("行情年龄（秒）", _text(status.get("market_data_age_seconds"))),
        ("tick_count", _text(status.get("tick_count"))),
        ("bar_count", _text(status.get("bar_count"))),
        ("首 tick Bar 启用", "是" if bool(status.get("first_tick_bar_enabled")) else "否"),
        ("首 tick Bar 已生成", "是" if bool(status.get("first_tick_bar_emitted")) else "否"),
        ("首 tick Bar 跳过原因", _text(status.get("first_tick_bar_skip_reason"))),
        ("行情问题", _text(status.get("market_issue"))),
    ]
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in market_rows:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = value

    document.add_heading("四、真实订单链", level=1)
    order_chain = status.get("order_chain") or []
    real_orders = [order for order in order_chain if str(order.get("track") or "").lower() == "real"]
    if real_orders:
        order_table = document.add_table(rows=1, cols=9)
        order_table.style = "Table Grid"
        headers = ["次序", "委托号", "合约", "方向", "开平", "价格", "数量", "时间", "状态"]
        for index, header in enumerate(headers):
            order_table.rows[0].cells[index].text = header
        for index, order in enumerate(real_orders, start=1):
            attempt = int(order.get("attempt") or 0)
            sequence = "首单" if attempt == 0 else f"追价 {attempt}"
            cells = order_table.add_row().cells
            values = [
                sequence,
                _text(order.get("order_id")),
                _text(order.get("symbol")),
                _direction(order.get("direction")),
                _offset(order.get("offset")),
                _text(order.get("price")),
                _text(order.get("volume")),
                _format_time(order.get("created_at")),
                _status_label(order.get("status")),
            ]
            for cell_index, value in enumerate(values):
                cells[cell_index].text = value
    else:
        document.add_paragraph("本轨道没有真实委托记录。")

    document.add_heading("五、真实成交或撤单证据", level=1)
    filled_real = [order for order in real_orders if str(order.get("status") or "").lower() == "filled"]
    cancelled_real = [
        order
        for order in real_orders
        if str(order.get("status") or "").lower() in {"cancelled", "canceled"}
    ]
    if filled_real:
        document.add_paragraph("存在真实成交回报，委托号：" + "、".join(_text(item.get("order_id")) for item in filled_real))
    if cancelled_real:
        document.add_paragraph("真实委托已撤单，委托号：" + "、".join(_text(item.get("order_id")) for item in cancelled_real))
    if not filled_real and not cancelled_real:
        document.add_paragraph("没有可确认的真实成交或撤单证据。")

    simulated_orders = [
        order for order in order_chain if str(order.get("track") or "").lower() == "simulated"
    ]
    if simulated_orders:
        document.add_heading("六、模拟订单与成交证据", level=1)
        document.add_paragraph("手续费与保证金：未计入（0）；模拟成交不代表券商真实成交。")
        sim_table = document.add_table(rows=1, cols=9)
        sim_table.style = "Table Grid"
        headers = ["次序", "委托号", "合约", "方向", "开平", "价格", "数量", "时间", "状态"]
        for index, header in enumerate(headers):
            sim_table.rows[0].cells[index].text = header
        for order in simulated_orders:
            attempt = int(order.get("attempt") or 0)
            sequence = "首单" if attempt == 0 else f"追价 {attempt}"
            cells = sim_table.add_row().cells
            values = [
                sequence,
                _text(order.get("order_id")),
                _text(order.get("symbol")),
                _direction(order.get("direction")),
                _offset(order.get("offset")),
                _text(order.get("price")),
                _text(order.get("volume")),
                _format_time(order.get("created_at")),
                _status_label(order.get("status")),
            ]
            for cell_index, value in enumerate(values):
                cells[cell_index].text = value
        document.add_paragraph(f"模拟持仓：{int(status.get('simulated_position_volume') or 0)} 手")

    document.add_heading("七、券商持仓与委托对账", level=1)
    reconcile_rows = [
        ("券商目标合约持仓", _text(status.get("broker_position_volume"))),
        ("券商活动委托", "、".join(map(str, status.get("broker_active_order_ids") or [])) or "无"),
        ("新鲜对账通过", "是" if bool(status.get("reconcile_ok")) else "否"),
        ("模拟持仓", _text(status.get("simulated_position_volume"))),
    ]
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in reconcile_rows:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = value

    document.add_heading("八、风控与异常记录", level=1)
    risk_rows = [
        ("滚动限频余量", _text(status.get("rate_limit_remaining"))),
        ("限频下次可报时间", _text(status.get("rate_limit_retry_after_seconds"))),
        ("真实持仓截止时间", _format_time(status.get("hold_deadline_at"))),
        ("失败原因", _text(status.get("failure_code"))),
        ("执行警告", _text(status.get("execution_warning"))),
    ]
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in risk_rows:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = value

    document.add_heading("九、最终结论", level=1)
    paragraph = document.add_paragraph(_conclusion(status))
    for run in paragraph.runs:
        run.font.size = Pt(12)
        run.bold = True

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
